import io
import re
from pathlib import Path

from sqlalchemy import delete, select

from app.config import storage_path
from app.models import Document, DocumentChunk
from app.services import audit

DOC_TYPE_KEYWORDS = [
    ("ITT", ["itt", "itb", "instruction", "tender document"]),
    ("GCC-PCC", ["gcc", "pcc", "conditions of contract", "general conditions"]),
    ("BOQ", ["boq", "bill of quantit", "price schedule", "priced"]),
    ("Specs", ["spec", "technical provision", "special provision"]),
    ("Forms", ["form", "appendices", "appendix"]),
    ("Drawings", ["drawing", "layout", "single line", "sld"]),
]

CHUNK_SIZE = 1200
CHUNK_OVERLAP = 150


def classify_doc(filename: str) -> str:
    f = filename.lower()
    for doc_type, keywords in DOC_TYPE_KEYWORDS:
        if any(k in f for k in keywords):
            return doc_type
    return "Other"


def extract_text(filename: str, content: bytes) -> str:
    name = filename.lower()
    try:
        if name.endswith(".pdf"):
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(content))
            return "\n".join((page.extract_text() or "") for page in reader.pages)[:400_000]
        if name.endswith(".docx"):
            from docx import Document as DocxDocument

            doc = DocxDocument(io.BytesIO(content))
            return "\n".join(p.text for p in doc.paragraphs)[:400_000]
        if name.endswith((".html", ".htm")):
            from bs4 import BeautifulSoup

            return BeautifulSoup(content.decode("utf-8", errors="ignore"), "html.parser").get_text(" ")[:400_000]
        return content.decode("utf-8", errors="ignore")[:400_000]
    except Exception as exc:
        return f"[PARSE_ERROR] {exc}"


def chunk_text(text: str, size: int = CHUNK_SIZE, overlap: int = CHUNK_OVERLAP):
    chunks = []
    start = 0
    while start < len(text):
        end = min(len(text), start + size)
        body = text[start:end].strip()
        if body:
            chunks.append(body)
        if end >= len(text):
            break
        start = end - overlap
    return chunks


def save_document(db, opp, filename: str, content: bytes) -> Document:
    folder = storage_path() / "documents" / str(opp.id)
    folder.mkdir(parents=True, exist_ok=True)
    safe_name = re.sub(r"[^A-Za-z0-9._\- ]", "_", filename)[:180]
    path = folder / safe_name
    path.write_bytes(content)

    text = extract_text(filename, content)
    doc = Document(
        opportunity_id=opp.id,
        filename=safe_name,
        doc_type=classify_doc(filename),
        file_path=str(path),
        size=len(content),
        parsed=not text.startswith("[PARSE_ERROR]"),
        parse_status="" if not text.startswith("[PARSE_ERROR]") else text[:200],
        content_text=text,
    )
    db.add(doc)
    db.commit()
    db.refresh(doc)

    db.execute(delete(DocumentChunk).where(DocumentChunk.document_id == doc.id))
    for i, chunk in enumerate(chunk_text(text)):
        db.add(DocumentChunk(document_id=doc.id, opportunity_id=opp.id, chunk_index=i, text=chunk, anchor=f"{doc.doc_type} #{i+1}"))
    db.commit()

    if opp.documents_status in ("none", "metadata_only"):
        opp.documents_status = "uploaded"
        db.commit()
    audit.record(db, "document.saved", opportunity_id=opp.id, detail={"filename": safe_name, "type": doc.doc_type, "size": len(content)})
    return doc


def fetch_for_opportunity(db, opp) -> int:
    from app.connectors.registry import get_connector
    from app.models import Source

    fetched = 0
    refs = list(opp.source_refs or [])
    for ref in refs:
        source_id = ref.get("source_id")
        if not source_id:
            # A converted lead can carry a source reference with no source id.
            # db.get() with a None key is a warning today and an error later.
            continue
        source = db.get(Source, source_id)
        if not source:
            continue
        connector = get_connector(source)
        if not connector:
            continue
        native_id = ref.get("native_id") or ""
        documents = connector.fetch_documents(native_id, ref=ref)
        if documents is None:
            continue
        for filename, content in documents:
            if len(content) > 50 * 1024 * 1024:
                continue
            save_document(db, opp, filename, content)
            fetched += 1
    if fetched == 0 and opp.documents_status == "none":
        fallback = (
            f"Notice summary (no attachments retrieved).\n"
            f"Title: {opp.title}\nBuyer: {opp.buyer_name}\nReference: {opp.buyer_ref}\n"
            f"Description: {opp.description}\n"
        )
        save_document(db, opp, "notice_summary.txt", fallback.encode("utf-8"))
        fetched = 1
    opp.documents_status = "fetched" if fetched else opp.documents_status
    db.commit()
    return fetched


def search_chunks(db, opportunity_id: int, query: str, limit: int = 5):
    terms = [t for t in re.findall(r"[a-z0-9]{3,}", query.lower())]
    chunks = db.execute(select(DocumentChunk).where(DocumentChunk.opportunity_id == opportunity_id)).scalars().all()
    scored = []
    for chunk in chunks:
        text_l = chunk.text.lower()
        score = sum(text_l.count(term) for term in terms)
        if score:
            scored.append((score, chunk))
    scored.sort(key=lambda x: -x[0])
    return [c for _, c in scored[:limit]]
